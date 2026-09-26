# -*- coding: utf-8 -*-
"""
cleanup_placeholders.py
Xoa bo hoan toan 4 doan ghi chu tam:
1. Doan 163 tai Muc 1.2 (va bo ' (VietinBank, n.d.)' o doan 162)
2. Doan 331 tai Muc 3.3.3
3. Doan 518 tai Muc 4.10
4. Doan 552 tai References (VietinBank n.d.)
Dong thoi chen hinh anh do hoa figure_4_1_empirical_model.png vao Hinh 4.1.
"""
import sys, os
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding='utf-8')

FILE = "workingfile/DTL_Master_Thesis_Draft.docx"
IMG_FILE = "workingfile/figure_4_1_empirical_model.png"

def main():
    doc = docx.Document(FILE)
    
    # 1. Muc 1.2: Bo ' (VietinBank, n.d.)' va xoa doan ghi chu so lieu
    for p in doc.paragraphs:
        if "(VietinBank, n.d.)" in p.text:
            p.text = p.text.replace(" (VietinBank, n.d.)", "")
            print("Da bo '(VietinBank, n.d.)' tai Muc 1.2")
            break

    paras_to_remove = []
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("[Cần bổ sung số liệu: quy mô dư nợ bảo lãnh"):
            paras_to_remove.append(p)
            print("Tim thay doan ghi chu Muc 1.2 can xoa")
        elif txt.startswith("[Cần bổ sung: căn cứ chấp thuận của VietinBank"):
            paras_to_remove.append(p)
            print("Tim thay doan ghi chu Muc 3.3.3 can xoa")
        elif txt.startswith("[Cần bổ sung sau khi có kết quả: các hạn chế"):
            paras_to_remove.append(p)
            print("Tim thay doan ghi chu Muc 4.10 can xoa")
        elif txt.startswith("VietinBank (n.d.) Annual Report and Audited Consolidated"):
            paras_to_remove.append(p)
            print("Tim thay doan VietinBank (n.d.) tai References can xoa")

    # Xoa cac paragraph
    for p in paras_to_remove:
        p._element.getparent().remove(p._element)

    # 2. Chen anh Figure 4.1
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip() == "Figure 4.1: Empirical regression results and standardized path coefficients":
            # Paragraph tiep theo la text mo hinh
            p_next = doc.paragraphs[i+1]
            if p_next.text.strip().startswith("[Empirical Path Model:"):
                p_next.text = ""
                p_next.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_next.paragraph_format.space_after = Pt(4)
                p_next.paragraph_format.line_spacing = 1.0
                r = p_next.add_run()
                r.add_picture(IMG_FILE, width=Inches(6.2))
                
                # Them source ngay duoi anh
                p_src = doc.add_paragraph()
                p_src.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_src.paragraph_format.space_after = Pt(12)
                p_src.paragraph_format.line_spacing = 1.5
                r_src = p_src.add_run("Source: Constructed from empirical regression estimates by the author (n = 800).")
                r_src.font.name = "Times New Roman"
                r_src.font.size = Pt(10)
                r_src.italic = True
                p_next._element.addnext(p_src._element)
                print("Da chen anh Figure 4.1 thanh cong!")
            break

    doc.save(FILE)
    print("Da luu file docx da don dep thanh cong.")

if __name__ == '__main__':
    main()
