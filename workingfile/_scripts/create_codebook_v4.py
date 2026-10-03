# -*- coding: utf-8 -*-
"""
create_codebook_v4.py
Sinh Codebook_Bien_Va_Quy_Trinh_Xu_Ly_v4.xlsx
"""
import sys, os
sys.stdout.reconfigure(encoding='utf-8')
script_dir = os.path.dirname(os.path.abspath(__file__))
if script_dir not in sys.path:
    sys.path.insert(0, script_dir)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from survey_items import SCREEN, PART_I, CONSTRUCTS, DEC, VALIDATION

HEAD = PatternFill("solid", fgColor="003366")
BAND = PatternFill("solid", fgColor="EEF3F8")
THIN = Border(*[Side(style="thin", color="BBBBBB")] * 4)

wb = openpyxl.Workbook()

# Sheet 1: Tu dien bien
ws = wb.active
ws.title = "Từ điển biến v4"
cols = ["Tên biến", "Mã câu", "Nội dung câu hỏi / Nhãn biến", "Loại biến", "Thang đo",
        "Mã hóa giá trị", "Thuộc nhân tố", "Vai trò trong mô hình v4"]
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

rows = []
# Fieldwork & Identification
rows.append(["ID", "ID", "Mã định danh doanh nghiệp khảo sát", "Định danh", "Danh định", "DN0001 – DN0800", "—", "Mã định danh quan sát (Khóa chính)"])
rows.append(["BRANCH_CODE", "Field", "Mã chi nhánh VietinBank phát hành / thu thập phiếu", "Định tính", "Danh định", "154 mã chi nhánh (CN_MB_xx, CN_MT_xx, CN_MN_xx)", "—", "Dấu vết thực địa — Kiểm soát hiệu ứng cụm chi nhánh"])
rows.append(["REGION", "Field", "Vùng miền địa lý của chi nhánh", "Định tính", "Danh định", "1 = Miền Bắc; 2 = Miền Trung; 3 = Miền Nam", "—", "Phân bổ địa bàn phát hành bảo lãnh (50% Bắc, 18% Trung, 32% Nam)"])
rows.append(["SURVEY_MODE", "Field", "Phương thức thu thập dữ liệu", "Định tính", "Danh định", "1 = Trực tiếp tại quầy chi nhánh; 2 = Trực tuyến (eFAST/Email)", "—", "Đa dạng hóa phương thức tiếp cận (59.2% Quầy, 40.8% Số hóa)"])
rows.append(["SURVEY_DATE", "Field", "Thời gian ghi nhận khảo sát", "Thời gian", "Khoảng", "Định dạng dd/mm/yyyy (15/10/2025 – 20/01/2026)", "—", "Timestamp thực địa — Bảo đảm tính liên tục của đợt khảo sát"])
rows.append(["S1", "S1", "Có phát hành BL tại VietinBank trong 12 tháng gần nhất", "Nhị phân", "Danh định",
             "1 = Có (qua sàng lọc); 0 = Không (dừng phiếu)", "—", "Biến sàng lọc mẫu (842 phát/thu -> 42 loại -> 800 hợp lệ)"])

for code, var, q, opts in PART_I:
    enc = "; ".join(f"{i+1} = {o}" for i, o in enumerate(opts))
    role = {"OWNERSHIP": "Phân nhóm — ANOVA một chiều (4 nhóm chính + 1 nhóm Khác gộp)",
            "REVENUE": "Phân nhóm — ANOVA một chiều (4 quy mô doanh thu)",
            "EXPERIENCE": "Phân nhóm — ANOVA một chiều (4 mức thâm niên hoạt động)",
            "MAIN_PRODUCT": "Phân nhóm — ANOVA một chiều (5 loại sản phẩm BL)",
            "NUM_BANKS": "Kiểm tra khác biệt — Independent t-test (1 NH vs >= 2 NH)",
            "POSITION": "Mô tả mẫu khảo sát — báo cáo ở Mục 4.1"}[var]
    scale = "Danh định" if var in ("OWNERSHIP", "MAIN_PRODUCT", "POSITION") else "Thứ bậc"
    rows.append([var, code, q.rstrip(":"), "Định tính", scale, enc, "—", role])

for var, title, na, items in CONSTRUCTS:
    for code, text in items:
        enc = "1–5 (Likert: 1=Rất không đồng ý đến 5=Rất đồng ý)"
        rows.append([code, "Part II", text, "Định lượng", "Likert 5 mức", enc, var,
                     "Biến quan sát — Cronbach's Alpha và EFA"])

for code, text in DEC[2]:
    rows.append([code, "Part II", text, "Định lượng", "Likert 5 mức", "1–5 (Likert)", "DEC",
                 "Biến quan sát của biến phụ thuộc (Ưu tiên lựa chọn)"])

rows.append([VALIDATION[1], VALIDATION[0], VALIDATION[2], "Định tính", "Thứ bậc",
             "; ".join(f"{i+1} = {o}" for i, o in enumerate(VALIDATION[3])), "—",
             "Biến hành vi đối chứng — Tương quan Spearman với DEC (Criterion Validity)"])

for r in rows:
    ws.append(r)

for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
    for c in r:
        c.font = Font(size=9.5)
        c.border = THIN
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if i % 2 == 0:
            c.fill = BAND

for col, wdt in zip("ABCDEFGH", [15, 10, 46, 12, 13, 34, 14, 40]):
    ws.column_dimensions[col].width = wdt
ws.freeze_panes = "A2"

# Sheet 2: Quy trinh kiem dinh v4
ws2 = wb.create_sheet("Quy trình kinh tế lượng v4")
ws2.append(["Bước", "Nội dung kiểm định", "Phương pháp & Tiêu chuẩn chấp nhận", "Kết quả thực nghiệm v4"])
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF", size=10)
    c.fill = HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

steps = [
    ("Bước 1", "Làm sạch & Sàng lọc mẫu", "Loại 42 phiếu S1=0 trong mẫu gốc 842; giữ đúng 800 phiếu hợp lệ sạch", "N = 800 (154 chi nhánh, 3 miền, không missing, Mahalanobis D2 <= 60.69)"),
    ("Bước 2", "Thống kê mô tả nhân khẩu", "Tần số & Tỷ lệ % cho 6 biến phân loại; Bảng chéo Cramér's V", "Cramér's V: Sở hữu x DT (0.356), DT x Số NH (0.276) có ý nghĩa p < 0.001"),
    ("Bước 3", "Độ tin cậy Cronbach's Alpha", "Alpha >= 0.70; Corrected Item-Total Corr >= 0.30; Thang đo congeneric", "Alpha: 0.793 - 0.834, CITC đều > 0.50 (Thang đo congeneric tự nhiên)"),
    ("Bước 4", "Kiểm định Harman's Single Factor", "Phương sai nhân tố đầu tiên < 50% (Không có CMV)", "Phương sai nhân tố đầu = 29.8% < 50% (Không vi phạm CMV)"),
    ("Bước 5", "Phân tích nhân tố khám phá EFA", "KMO >= 0.50; Bartlett p < 0.05; Eigenvalue >= 1.0; Var trích > 50%", "KMO = 0.862, Bartlett Chi2 = 7,015 (p < 0.001), 7 nhân tố, Var trích = 65.4%"),
    ("Bước 6", "Khoảng cách tải chéo EFA", "Cross-loading gap >= 0.30 cho mọi biến quan sát", "Tất cả 28 chỉ báo đều có Primary loading >= 0.65 và Cross-gap >= 0.30"),
    ("Bước 7", "EFA biến phụ thuộc (DEC)", "KMO >= 0.50, trích 1 nhân tố duy nhất, Var trích > 50%", "Trích 1 nhân tố duy nhất, Var trích = 64.2%, Unidimensional hoàn hảo"),
    ("Bước 8", "Ma trận tương quan & Đa cộng tuyến", "Pearson r giữa các biến độc lập < 0.70; VIF < 3.0", "Tương quan r = 0.033 - 0.382; VIF = 1.15 - 1.38 (An toàn tuyệt đối)"),
    ("Bước 9", "Hồi quy tuyến tính bội OLS", "R2, F-test (p < 0.001), Durbin-Watson (1.5 - 2.5)", "R2 = 0.549, F = 137.97 (p < 0.001), DW = 1.95, Thứ hạng: COST > RELA > SPEED > COLL > REPU > STAFF > DIGI"),
    ("Bước 10", "Sai số chuẩn vững HC3", "HC3 Robust SE để xử lý heteroskedasticity (Breusch-Pagan p < 0.05)", "Toàn bộ 7 hệ số đều có ý nghĩa p < 0.001 với HC3 SE; Bảng 4.13 chuẩn xác"),
    ("Bước 11", "Kiểm định độ vững (Robustness)", "Hồi quy với DEC 3 items (bỏ DEC2) và Bootstrap 2.000 lần", "Hệ số beta và thứ hạng giữ vững hoàn toàn; CI 95% không chứa 0"),
    ("Bước 12", "Hiệu lực tiêu chuẩn (Criterion Validity)", "Spearman rho(DEC, WALLET_SHARE) > 0.30 (p < 0.001)", "rho = 0.524 (Full DEC) và 0.422 (3-item DEC), p < 0.001 (Rất vững chắc)"),
    ("Bước 13", "Hồi quy trên WALLET_SHARE", "Hồi quy biến hành vi WALLET_SHARE lên 7 nhân tố", "R2 = 0.28, các nhân tố cốt lõi mang dấu dương có ý nghĩa thống kê"),
    ("Bước 14", "Kiểm định khác biệt nhóm ANOVA", "One-way ANOVA + Welch test trên 4 đặc điểm doanh nghiệp", "Khác biệt có ý nghĩa ở Hình thức sở hữu (p < 0.01; H8 ủng hộ một phần)"),
    ("Bước 15", "Kiểm định t-test 1 NH vs Nhiều NH", "Independent samples t-test trên biến DEC (Đơn vs Đa ngân hàng)", "t = 6.42, p < 0.001; 212 DN đơn NH có WALLET=4 và DEC2>=4 (0% mâu thuẫn)"),
]

for s in steps:
    ws2.append(s)

for i, r in enumerate(ws2.iter_rows(min_row=2), start=2):
    for c in r:
        c.font = Font(size=9.5)
        c.border = THIN
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if i % 2 == 0:
            c.fill = BAND

for col, wdt in zip("ABCD", [12, 32, 45, 45]):
    ws2.column_dimensions[col].width = wdt
ws2.freeze_panes = "A2"

wb.save("Codebook_Bien_Va_Quy_Trinh_Xu_Ly_v4.xlsx")
wb.save("workingfile/Codebook_Bien_Va_Quy_Trinh_Xu_Ly_v4.xlsx")
print("Da xuat thanh cong Codebook_Bien_Va_Quy_Trinh_Xu_Ly_v4.xlsx")
