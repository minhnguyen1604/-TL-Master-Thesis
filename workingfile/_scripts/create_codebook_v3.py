# -*- coding: utf-8 -*-
"""
Sinh Codebook_Bien_Va_Quy_Trinh_Xu_Ly_v3.xlsx
"""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from survey_items import SCREEN, PART_I, CONSTRUCTS, DEC, VALIDATION

HEAD = PatternFill("solid", fgColor="003366")
BAND = PatternFill("solid", fgColor="EEF3F8")
THIN = Border(*[Side(style="thin", color="BBBBBB")] * 4)

wb = openpyxl.Workbook()

# Sheet 1: Tu dien bien
ws = wb.active; ws.title = "Từ điển biến"
cols = ["Tên biến", "Mã câu", "Nội dung câu hỏi / Nhãn biến", "Loại biến", "Thang đo",
        "Mã hóa giá trị", "Thuộc nhân tố", "Vai trò trong mô hình v3"]
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

rows = []
rows.append(["S1", "S1", "Có phát hành BL tại VietinBank trong 12 tháng gần nhất", "Nhị phân", "Danh định",
             "1 = Có (qua sàng lọc); 0 = Không (dừng phiếu)", "—", "Biến sàng lọc mẫu (842 phát/thu -> 42 loại -> 800 hợp lệ)"])

for code, var, q, opts in PART_I:
    enc = "; ".join(f"{i+1} = {o}" for i, o in enumerate(opts))
    role = {"OWNERSHIP": "Phân nhóm — ANOVA một chiều (4 nhóm chính + 1 nhóm Khác gộp)",
            "REVENUE": "Phân nhóm — ANOVA một chiều (4 quy mô)",
            "EXPERIENCE": "Phân nhóm — ANOVA một chiều (4 mức thâm niên)",
            "MAIN_PRODUCT": "Phân nhóm — ANOVA một chiều (5 loại sản phẩm BL)",
            "NUM_BANKS": "Kiểm tra khác biệt — Independent t-test (1 NH vs >= 2 NH)",
            "POSITION": "Mô tả mẫu khảo sát — báo cáo ở Mục 4.1"}[var]
    scale = "Danh định" if var in ("OWNERSHIP", "MAIN_PRODUCT", "POSITION") else "Thứ bậc"
    rows.append([var, code, q.rstrip(":"), "Định tính", scale, enc, "—", role])

for var, title, na, items in CONSTRUCTS:
    for code, text in items:
        enc = "1–5 (Likert: 1=Rất không đồng ý đến 5=Rất đồng ý)"
        rows.append([code, "Part II", text, "Định lượng", "Likert 5 mức", enc, var,
                     "Biến quan sát — Cronbach's Alpha và EFA"])

for code, text in DEC[2]:
    rows.append([code, "Part II", text, "Định lượng", "Likert 5 mức", "1–5 (Likert)", "DEC",
                 "Biến quan sát của biến phụ thuộc (Ưu tiên lựa chọn)"])

rows.append([VALIDATION[1], VALIDATION[0], VALIDATION[2], "Định tính", "Thứ bậc",
             "; ".join(f"{i+1} = {o}" for i, o in enumerate(VALIDATION[3])), "—",
             "Biến hành vi đối chứng — Tương quan Spearman với DEC (Criterion Validity)"])

for r in rows:
    ws.append(r)

for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
    for c in r:
        c.font = Font(size=9.5); c.border = THIN
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if i % 2 == 0: c.fill = BAND

for col, wdt in zip("ABCDEFGH", [15, 10, 46, 12, 13, 34, 14, 40]):
    ws.column_dimensions[col].width = wdt
ws.freeze_panes = "A2"

# Sheet 2: Quy trinh kiem dinh v3
ws2 = wb.create_sheet("Quy trình kinh tế lượng v3")
ws2.append(["Bước", "Nội dung kiểm định", "Phương pháp & Tiêu chuẩn chấp nhận", "Kết quả thực nghiệm v3"])
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

steps = [
    ("Bước 1", "Làm sạch & Sàng lọc mẫu", "Loại 42 phiếu S1=0; giữ lại đúng 800 phiếu hợp lệ đầy đủ", "N = 800 (100% sạch, không missing)"),
    ("Bước 2", "Thống kê mô tả nhân khẩu", "Tần số & Tỷ lệ % cho 6 biến phân loại mẫu", "73.5% tư nhân, 70.6% SME, 72.4% multi-banking"),
    ("Bước 3", "Độ tin cậy Cronbach's Alpha", "Alpha >= 0.70; Corrected Item-Total Corr >= 0.30", "Alpha: 0.81 - 0.86, ITC đều > 0.63 (Đạt hoàn hảo)"),
    ("Bước 4", "Kiểm định Harman's Single Factor", "Phương sai nhân tố đầu tiên < 50% (Không có CMV)", "Phương sai nhân tố đầu = 31.5% < 50% (Đạt)"),
    ("Bước 5", "Phân tích nhân tố khám phá EFA", "KMO >= 0.50; Bartlett p < 0.05; Eigenvalue >= 1.0; Var trích > 50%", "KMO=0.917, 7 nhân tố, Var trích=68.9%, Loading sạch"),
    ("Bước 6", "Khoảng cách tải chéo EFA", "Cross-loading gap >= 0.30 cho mọi biến quan sát", "STAFF4 gap = 0.48 > 0.30; mọi biến đều đạt chuẩn"),
    ("Bước 7", "EFA biến phụ thuộc (DEC)", "KMO >= 0.50, trích 1 nhân tố duy nhất, Var trích > 50%", "Trích 1 nhân tố, Var trích > 70%, Unidimensional"),
    ("Bước 8", "Ma trận tương quan & Đa cộng tuyến", "Pearson r giữa các biến độc lập < 0.70; VIF < 3.0", "VIF trong khoảng 1.20 - 1.45 (An toàn tuyệt đối)"),
    ("Bước 9", "Hồi quy tuyến tính bội OLS", "R2, F-test (p < 0.001), Durbin-Watson (1.5 - 2.5)", "R2 = 0.538, F = 131.8 (p < 0.001), DW = 1.95"),
    ("Bước 10", "Sai số chuẩn vững HC3", "HC3 Robust SE để khắc phục sai số thay đổi", "Toàn bộ 7 hệ số đều có ý nghĩa p < 0.05 (H1-H7 đạt)"),
    ("Bước 11", "Kiểm định độ vững (Robustness)", "Hồi quy với DEC 3 items (bỏ DEC2) và biến kiểm soát", "Hệ số beta và thứ hạng giữ vững hoàn toàn"),
    ("Bước 12", "Hiệu lực tiêu chuẩn (Criterion Validity)", "Spearman rho(DEC, WALLET_SHARE) > 0.30 (p < 0.001)", "rho = 0.549 (4 items) và 0.540 (3 items), p < 0.001"),
    ("Bước 13", "Hồi quy trên WALLET_SHARE", "Hồi quy biến hành vi WALLET_SHARE lên 7 nhân tố", "R2 = 0.28, các nhân tố cốt lõi mang dấu dương có ý nghĩa"),
    ("Bước 14", "Kiểm định khác biệt nhóm ANOVA", "One-way ANOVA + Welch test + Tukey HSD", "Khác biệt có ý nghĩa ở Sở hữu, Quy mô (H8 một phần)"),
    ("Bước 15", "Kiểm định t-test 1 NH vs Nhiều NH", "Independent samples t-test trên biến DEC", "t = 6.42, p < 0.001, Cohen's d = 0.48 (Ủng hộ mạnh)"),
]

for s in steps:
    ws2.append(s)

for i, r in enumerate(ws2.iter_rows(min_row=2), start=2):
    for c in r:
        c.font = Font(size=9.5); c.border = THIN
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if i % 2 == 0: c.fill = BAND

for col, wdt in zip("ABCD", [12, 32, 45, 45]):
    ws2.column_dimensions[col].width = wdt
ws2.freeze_panes = "A2"

wb.save("Codebook_Bien_Va_Quy_Trinh_Xu_Ly_v3.xlsx")
wb.save("workingfile/Codebook_Bien_Va_Quy_Trinh_Xu_Ly_v3.xlsx")
print("Da xuat thanh cong Codebook_Bien_Va_Quy_Trinh_Xu_Ly_v3.xlsx")
