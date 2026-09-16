# -*- coding: utf-8 -*-
"""Cap nhat ban thao theo bo cau hoi 32 cau."""
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from survey_items import CONSTRUCTS, DEC

FILE = "DTL_Master_Thesis_Draft.docx"

EN = {
 "COMP1": "The guarantee issuance fee charged by VietinBank is reasonable relative to the service quality received.",
 "COMP2": "VietinBank's guarantee fee schedule is competitive compared with that of other banks.",
 "COMP3": "VietinBank offers fee reductions to enterprises that transact regularly.",
 "COMP4": "Ancillary charges (amendment, extension, enquiry) at VietinBank are at an acceptable level.",
 "SPEED1": "Appraisal and approval of the guarantee limit at VietinBank is fast.",
 "SPEED2": "The application file required by VietinBank is concise and does not demand excessive documentation.",
 "SPEED3": "The time from submission of a complete file to receipt of the letter of guarantee meets contract deadlines.",
 "SPEED4": "Amendments and extensions to a letter of guarantee are processed quickly by VietinBank.",
 "DIGI1": "The enterprise can submit guarantee applications online through the VietinBank eFAST platform.",
 "DIGI2": "VietinBank issues digitally signed electronic letters of guarantee promptly.",
 "DIGI3": "Checking the status of a guarantee application online is convenient.",
 "DIGI4": "Beneficiaries accept VietinBank's electronic letter of guarantee without requiring a paper original.",
 "REPU1": "VietinBank's reputation ranks among the leaders of the Vietnamese banking market.",
 "REPU2": "Letters of guarantee issued by VietinBank are widely accepted by project owners and procuring entities.",
 "REPU3": "VietinBank's financial strength adds credibility to the enterprise when bidding.",
 "REPU4": "The enterprise is confident that VietinBank would honour its obligation if a guarantee were called.",
 "RELA1": "A long-standing credit relationship with VietinBank makes it easier to obtain a guarantee.",
 "RELA2": "Using other VietinBank services brings advantages for the enterprise's guarantee transactions.",
 "RELA3": "VietinBank grants a guarantee limit appropriate to the enterprise's needs.",
 "RELA4": "VietinBank adjusts the guarantee limit promptly when the enterprise's needs change.",
 "STAFF1": "VietinBank officers have solid technical expertise in guarantee operations.",
 "STAFF2": "VietinBank officers can advise on guarantee wording consistent with the legislation in force.",
 "STAFF3": "VietinBank officers warn the enterprise of risks in the guarantee terms before issuance.",
 "STAFF4": "VietinBank officers respond promptly when the enterprise raises a problem.",
 "COLL1": "VietinBank applies a cash margin ratio appropriate to the enterprise's capacity.",
 "COLL2": "VietinBank reduces or waives the cash margin for qualifying clients.",
 "COLL3": "VietinBank accepts a diverse range of collateral.",
 "COLL4": "Valuation and registration of collateral at VietinBank is quick.",
 "DEC1": "When guarantee needs arise, the enterprise prioritises VietinBank ahead of other banks.",
 "DEC2": "The enterprise places the majority of its guarantee value with VietinBank rather than with other banks.",
 "DEC3": "For upcoming tenders and contracts, the enterprise intends to continue with VietinBank.",
 "DEC4": "The enterprise would recommend VietinBank to partners as the best issuing bank.",
}
NAMES = {"COST_COMP": "Price Competitiveness", "PROC_SPEED": "Processing Speed",
         "DIGITAL_CONV": "Digital eFAST Convenience", "BANK_REP": "Bank Reputation",
         "RELATIONSHIP": "Relationship & Limits", "STAFF_QUAL": "Staff Professionalism",
         "COLL_POLICY": "Collateral & Margin Flexibility"}

TEXT = [
 ("Verifying that the three indicators of each construct", "Verifying that the four indicators of each construct"),
 ("Each construct is measured by three indicators, with the exception of the dependent construct, which "
  "is measured by four. Three indicators is the minimum required for a construct to be identified in "
  "factor analysis and is adopted deliberately here to keep the instrument short.",
  "Each construct, including the dependent construct, is measured by four indicators. Three indicators is "
  "the minimum required for a construct to be identified in factor analysis; a fourth is included as "
  "headroom, so that if an indicator is removed at the reliability or factor analysis stage the construct "
  "still retains three and need not be abandoned."),
 ("3.2.2. Mapping Scales with the Official 25-Item Survey Questionnaire",
  "3.2.2. Mapping Scales with the Official 32-Item Survey Questionnaire"),
 ("Part II contains the twenty-five Likert indicators listed in",
  "Part II contains the thirty-two Likert indicators listed in"),
 ("which for twenty-five indicators yields 125 responses", "which for thirty-two indicators yields 160 responses"),
 ("For each of the twenty-five indicators the analysis reports", "For each of the thirty-two indicators the analysis reports"),
 ("Because each construct carries only three indicators and two are the "
  "minimum for identification, removal is undertaken only where an indicator is clearly deficient on "
  "both criteria, and any such decision is reported explicitly with its justification rather than "
  "applied silently.",
  "Because each construct carries four indicators, the removal of one still leaves three, which is "
  "sufficient for identification; removal is nonetheless undertaken only where an indicator is clearly "
  "deficient on both criteria, and any such decision is reported explicitly with its justification."),
 ("examines whether the twenty-five indicators group empirically into the eight",
  "examines whether the thirty-two indicators group empirically into the eight"),
 ("separately for the twenty-one independent indicators and the four dependent",
  "separately for the twenty-eight independent indicators and the four dependent"),
 ("Data screening removes three categories of response before analysis.",
  "Data screening removes four categories of response before analysis. Questionnaires failing the "
  "screening question, that is those reporting no guarantee issued at VietinBank in the preceding twelve "
  "months, are excluded at the outset."),
 ("The survey instrument is organised in two parts. Part I collects six items of classification "
  "information:",
  "The survey instrument is organised in four parts. Part A is a single screening question establishing "
  "that the enterprise has had a guarantee issued at VietinBank within the preceding twelve months; "
  "enterprises answering no do not proceed. Part B collects six items of classification information:"),
]


def shade(cell, f="D9D9D9"):
    tcPr = cell._element.get_or_add_tcPr()
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), f); tcPr.append(s)


def settext(p, t):
    if not p.runs: p.add_run(t); return
    p.runs[0].text = t
    for r in p.runs[1:]: r.text = ""


def main():
    d = docx.Document(FILE)

    # ---- 1. thay van ban
    hit = 0
    for old, new in TEXT:
        for p in d.paragraphs:
            if old in p.text:
                settext(p, p.text.replace(old, new)); hit += 1; break
    print(f"Van ban: thay duoc {hit}/{len(TEXT)}")

    # ---- 2. dung lai Bang 3.1
    tgt = None
    for t in d.tables:
        if t.rows and t.rows[0].cells[0].text.strip() == "Construct" and len(t.rows) > 20:
            tgt = t; break
    assert tgt is not None, "khong tim thay Bang 3.1"
    while len(tgt.rows) > 1:
        tgt._element.remove(tgt.rows[-1]._element)
    blocks = [(v, NAMES[v], items) for v, _, _, items in CONSTRUCTS] + \
             [("DEC", "Selection Priority & Patronage Intention", DEC[2])]
    W = [1.35, 0.65, 4.2]
    for var, name, items in blocks:
        for n, (code, _) in enumerate(items):
            cells = tgt.add_row().cells
            vals = [f"{var}\n({name})" if n == 0 else "", code, EN[code]]
            for k, v in enumerate(vals):
                cells[k].width = Inches(W[k]); cells[k].text = ""
                p = cells[k].paragraphs[0]
                p.paragraph_format.space_after = Pt(2); p.paragraph_format.line_spacing = 1.0
                r = p.add_run(v); r.font.name = 'Times New Roman'; r.font.size = Pt(9)
    print(f"Bang 3.1: dung lai voi {len(tgt.rows)-1} bien quan sat")

    # ---- 3. cap nhat dong nguon duoi Bang 3.1
    for p in d.paragraphs:
        if p.text.strip().startswith("Source: compiled by the author from the theoretical framework"):
            settext(p, "Source: compiled by the author from the theoretical framework in Chapter 2. "
                       "Total 32 Likert indicators = 8 constructs × 4 indicators.")
            break

    # ---- 4. them doan ve cau doi chung vao cuoi 3.2.2
    for i, p in enumerate(d.paragraphs):
        if p.text.strip().startswith("The dependent construct block is introduced by an explicit comparative"):
            el = OxmlElement('w:p'); p._element.addnext(el)
            np = docx.text.paragraph.Paragraph(el, p._parent)
            np.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            np.paragraph_format.space_after = Pt(8)
            np.paragraph_format.line_spacing = 1.5
            np.paragraph_format.first_line_indent = Inches(0.25)
            r = np.add_run(
                "Part D adds one behavioural item asking the enterprise to estimate the share of its total "
                "guarantee value placed with VietinBank over the preceding twelve months. This item is not a "
                "predictor and does not enter the regression. It serves as a criterion variable against which "
                "the attitudinal dependent scale is validated: a scale claiming to measure patronage priority "
                "should correlate with the share of business actually placed, and the association is tested by "
                "Spearman rank correlation and reported alongside the reliability results in Section 4.2. A "
                "further design feature concerns the digital block, where respondents who have not used the "
                "eFAST channel select a \"not used\" option rather than a scale point; treating such responses "
                "as missing prevents enterprises without experience of the channel from contributing guesses to "
                "the DIGITAL_CONV score, and the proportion selecting it is itself reported as a descriptive "
                "finding.")
            r.font.name = 'Times New Roman'; r.font.size = Pt(12)
            print("Da them doan ve cau doi chung va o 'chua su dung'")
            break

    d.save(FILE)


if __name__ == "__main__":
    main()
