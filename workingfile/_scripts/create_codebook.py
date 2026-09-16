# -*- coding: utf-8 -*-
"""Sinh Codebook.xlsx - tu dien bien phuc vu nhap lieu va phan tich."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from survey_items import SCREEN, PART_I, CONSTRUCTS, DEC, VALIDATION, N_ITEMS

HEAD = PatternFill("solid", fgColor="003366")
BAND = PatternFill("solid", fgColor="EEF3F8")
THIN = Border(*[Side(style="thin", color="BBBBBB")] * 4)

wb = openpyxl.Workbook()

# ---------------- Sheet 1: tu dien bien
ws = wb.active; ws.title = "Từ điển biến"
cols = ["Tên biến", "Câu hỏi", "Nhãn biến", "Loại biến", "Thang đo",
        "Mã hóa giá trị", "Thuộc nhân tố", "Vai trò trong mô hình"]
ws.append(cols)
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

rows = []
rows.append(["SCREEN", SCREEN[0], "Có phát hành BL tại VietinBank trong 12 tháng", "Nhị phân", "Danh định",
             "1 = Có; 2 = Không", "—", "Sàng lọc — loại toàn bộ phiếu có giá trị 2"])
for code, var, q, opts in PART_I:
    enc = "; ".join(f"{i+1} = {o}" for i, o in enumerate(opts))
    role = {"OWNERSHIP": "Phân nhóm — ANOVA một chiều + Tukey HSD",
            "REVENUE": "Phân nhóm — ANOVA một chiều + Tukey HSD",
            "EXPERIENCE": "Phân nhóm — ANOVA một chiều + Tukey HSD",
            "MAIN_PRODUCT": "Phân nhóm — ANOVA một chiều + Tukey HSD",
            "NUM_BANKS": "Kiểm soát — gộp nhị phân (1 vs 2,3,4) để chạy t-test",
            "POSITION": "Mô tả mẫu — chỉ báo cáo ở mục 4.1"}[var]
    scale = "Danh định" if var in ("OWNERSHIP", "MAIN_PRODUCT", "POSITION") else "Thứ bậc"
    rows.append([var, code, q.rstrip(":"), "Định tính", scale, enc, "—", role])
for var, title, na, items in CONSTRUCTS:
    for code, text in items:
        enc = "1–5 (Likert)" + ("; 9 = Chưa sử dụng (coi là khuyết)" if na else "")
        rows.append([code, "Phần C", text, "Định lượng", "Likert 5 mức", enc, var,
                     "Biến quan sát — vào Cronbach's Alpha và EFA"])
for code, text in DEC[2]:
    rows.append([code, "Phần C", text, "Định lượng", "Likert 5 mức", "1–5 (Likert)", "DEC",
                 "Biến quan sát của biến phụ thuộc"])
rows.append([VALIDATION[1], VALIDATION[0], VALIDATION[2], "Định tính", "Thứ bậc",
             "; ".join(f"{i+1} = {o}" for i, o in enumerate(VALIDATION[3])), "—",
             "Biến đối chứng — tương quan Spearman với DEC để kiểm định giá trị tiêu chuẩn"])
for r in rows:
    ws.append(r)
for i, r in enumerate(ws.iter_rows(min_row=2), start=2):
    for c in r:
        c.font = Font(size=9.5); c.border = THIN
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if i % 2 == 0: c.fill = BAND
for col, wdt in zip("ABCDEFGH", [15, 10, 46, 12, 13, 34, 14, 40]):
    ws.column_dimensions[col].width = wdt
ws.freeze_panes = "A2"

# ---------------- Sheet 2: bien tong hop
ws2 = wb.create_sheet("Biến tổng hợp")
ws2.append(["Biến mô hình", "Tên đầy đủ", "Các biến quan sát", "Số câu", "Vai trò", "Dấu kỳ vọng", "Cách tính"])
for c in ws2[1]:
    c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
FULL = {"COST_COMP": "Price Competitiveness", "PROC_SPEED": "Processing Speed",
        "DIGITAL_CONV": "Digital eFAST Convenience", "BANK_REP": "Bank Reputation",
        "RELATIONSHIP": "Relationship & Limits", "STAFF_QUAL": "Staff Professionalism",
        "COLL_POLICY": "Collateral & Margin Flexibility"}
for n, (var, title, na, items) in enumerate(CONSTRUCTS, start=1):
    ws2.append([var, FULL[var], f"{items[0][0]}–{items[-1][0]}", len(items), f"Độc lập (X{n})", "+",
                "Trung bình cộng các biến quan sát CÒN LẠI SAU EFA"])
ws2.append(["DEC", "Selection Priority & Patronage Intention", "DEC1–DEC4", 4, "Phụ thuộc (Y)", "—",
            "Trung bình cộng các biến quan sát CÒN LẠI SAU EFA"])
for i, r in enumerate(ws2.iter_rows(min_row=2), start=2):
    for c in r:
        c.font = Font(size=10); c.border = THIN
        c.alignment = Alignment(vertical="center", wrap_text=True)
        if i % 2 == 0: c.fill = BAND
for col, wdt in zip("ABCDEFG", [16, 34, 16, 8, 14, 12, 45]):
    ws2.column_dimensions[col].width = wdt

# ---------------- Sheet 3: quy trinh xu ly
ws3 = wb.create_sheet("Quy trình xử lý")
ws3.append(["Bước", "Thao tác", "Ngưỡng / Quy tắc", "Ghi chú"])
for c in ws3[1]:
    c.font = Font(bold=True, color="FFFFFF", size=10); c.fill = HEAD
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
steps = [
 ["1", "Loại phiếu không đạt sàng lọc", "S1 = 2 (Không)", "Ghi lại số phiếu bị loại để báo cáo ở mục 4.1"],
 ["2", "Loại phiếu thiếu dữ liệu", "Thiếu bất kỳ câu nào ở Phần C", "Listwise deletion, không thay thế giá trị"],
 ["3", "Xử lý mã 9 của nhóm DIGI", "Coi 'Chưa sử dụng' là khuyết",
  "Báo cáo riêng tỷ lệ doanh nghiệp chưa dùng eFAST — đây là phát hiện có giá trị"],
 ["4", "Loại phiếu đánh lụi", "Cùng một mức cho toàn bộ 32 câu", "Straight-lining"],
 ["5", "Đảo mã câu nghịch đảo", "Nếu có câu nghịch đảo", "PHẢI làm trước khi tính trung bình"],
 ["6", "Cronbach's Alpha từng biến", "α ≥ 0,70; tương quan biến-tổng ≥ 0,30",
  "Loại câu không đạt; mỗi biến phải còn ≥ 3 câu"],
 ["7", "EFA biến độc lập", "KMO ≥ 0,50; Bartlett p < 0,05; eigenvalue ≥ 1; hệ số tải ≥ 0,50; chênh tải chéo ≥ 0,30",
  "Chạy riêng 28 câu độc lập"],
 ["8", "EFA biến phụ thuộc", "Kỳ vọng rút trích 1 nhân tố duy nhất", "Chạy riêng 4 câu DEC"],
 ["9", "Tính điểm biến tổng hợp", "Trung bình cộng các câu còn lại", "CHỈ thực hiện sau bước 7 và 8"],
 ["10", "Tương quan Pearson", "r giữa các biến độc lập < 0,80", "Kiểm tra sơ bộ đa cộng tuyến"],
 ["11", "Kiểm định VIF", "VIF < 3,0 (tolerance > 0,33)", "Ngưỡng chặt hơn mức 10 thông thường"],
 ["12", "Hồi quy OLS", "Kiểm tra 4 giả định; báo cáo beta chuẩn hóa", "Beta chuẩn hóa dùng để xếp hạng nhân tố"],
 ["13", "Kiểm định độ vững", "Thêm biến giả OWNERSHIP, REVENUE, NUM_BANKS", "Xác nhận β1–β7 ổn định"],
 ["14", "ANOVA / t-test", "Levene; nếu vi phạm dùng Welch; hậu kiểm Tukey HSD",
  "Nhóm có n < 30 phải gộp trước khi chạy"],
 ["15", "Kiểm định giá trị tiêu chuẩn", "Spearman giữa DEC và WALLET_SHARE",
  "Tương quan dương có ý nghĩa = bằng chứng thang đo đo đúng hành vi"],
]
for s in steps: ws3.append(s)
for i, r in enumerate(ws3.iter_rows(min_row=2), start=2):
    for c in r:
        c.font = Font(size=10); c.border = THIN
        c.alignment = Alignment(vertical="top", wrap_text=True)
        if i % 2 == 0: c.fill = BAND
for col, wdt in zip("ABCD", [7, 34, 52, 52]):
    ws3.column_dimensions[col].width = wdt

wb.save("Codebook_Bien_Va_Quy_Trinh_Xu_Ly.xlsx")
print("Da tao Codebook_Bien_Va_Quy_Trinh_Xu_Ly.xlsx |", len(rows), "bien |", len(steps), "buoc xu ly")
