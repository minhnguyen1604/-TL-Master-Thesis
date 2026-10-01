# -*- coding: utf-8 -*-
"""
insert_missing_references.py
Adds missing references (Kaiser 1974, Spence 1973, Law on Credit Institutions 2024, Decree 80/2021)
into the REFERENCES section of DTL_Master_Thesis_Draft.docx.
"""
import sys, os
import docx
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement

sys.stdout.reconfigure(encoding='utf-8')

FILE = "workingfile/DTL_Master_Thesis_Draft.docx"

def main():
    doc = docx.Document(FILE)

    new_references = [
        "Kaiser, H.F. (1974) 'An index of factorial simplicity', Psychometrika, 39(1), pp. 31–36.",
        "National Assembly of Vietnam (2024) Law on Credit Institutions No. 32/2024/QH15. Hanoi: National Assembly of the Socialist Republic of Vietnam.",
        "Spence, M. (1973) 'Job market signaling', The Quarterly Journal of Economics, 87(3), pp. 355–374.",
        "Vietnamese Government (2021) Decree No. 80/2021/ND-CP on elaboration of some articles of the Law on Supporting Small and Medium-Sized Enterprises. Hanoi: Government of the Socialist Republic of Vietnam."
    ]

    in_ref = False
    last_ref_p = None
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt == 'REFERENCES':
            in_ref = True
            continue
        if in_ref and (txt.startswith('Appendix') or txt.startswith('CHAPTER')):
            break
        if in_ref and txt:
            last_ref_p = p

    for ref_text in new_references:
        el = OxmlElement('w:p')
        last_ref_p._element.addnext(el)
        p = docx.text.paragraph.Paragraph(el, last_ref_p._parent)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        pf = p.paragraph_format
        pf.space_after = Pt(6)
        pf.line_spacing = 1.5
        pf.left_indent = Inches(0.5)
        pf.first_line_indent = Inches(-0.5)
        r = p.add_run(ref_text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(11)
        last_ref_p = p
        print(f"Added reference: {ref_text[:40]}...")

    doc.save(FILE)
    print("Successfully added all 4 missing references to DTL_Master_Thesis_Draft.docx!")

if __name__ == '__main__':
    main()
