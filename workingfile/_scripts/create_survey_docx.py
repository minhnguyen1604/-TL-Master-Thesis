# -*- coding: utf-8 -*-
"""Sinh file Word Phieu Khao Sat Chinh Thuc - 32 cau Likert, 7 bien x 4 cau + DEC 4 cau."""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from survey_items import SCREEN, PART_I, CONSTRUCTS, DEC, VALIDATION, N_ITEMS

NAVY = RGBColor(0x00, 0x33, 0x66)
GREY, LIGHT = "D9D9D9", "F2F2F2"
C, J = WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.JUSTIFY


def shade(cell, f):
    tcPr = cell._element.get_or_add_tcPr()
    s = OxmlElement('w:shd'); s.set(qn('w:val'), 'clear')
    s.set(qn('w:color'), 'auto'); s.set(qn('w:fill'), f); tcPr.append(s)


def P(doc, t, size=11.5, bold=False, italic=False, align=None, color=None, after=6):
    p = doc.add_paragraph()
    if align is not None: p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    r = p.add_run(t); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    return p


def build():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(0.9)
        s.left_margin = s.right_margin = Inches(0.9)
    st = doc.styles['Normal']; st.font.name = 'Times New Roman'; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    P(doc, "NGÂN HÀNG TMCP CÔNG THƯƠNG VIỆT NAM (VIETINBANK)", 11.5, True, align=C, color=NAVY, after=0)
    P(doc, "KHẢO SÁT Ý KIẾN KHÁCH HÀNG DOANH NGHIỆP VỀ DỊCH VỤ BẢO LÃNH", 10.5, True, True, align=C, color=NAVY)
    P(doc, "PHIẾU KHẢO SÁT CHÍNH THỨC", 15, True, align=C, color=NAVY, after=2)
    P(doc, "Các yếu tố ảnh hưởng đến quyết định lựa chọn dịch vụ bảo lãnh ngân hàng "
           "của khách hàng doanh nghiệp tại VietinBank", 11, False, True, align=C, after=10)

    P(doc, "Kính gửi Quý Doanh nghiệp,", 11, True, after=3)
    P(doc, f"Nhằm nâng cao chất lượng dịch vụ bảo lãnh, chúng tôi tiến hành khảo sát này. Phiếu gồm 01 câu sàng "
           f"lọc, 06 câu thông tin chung, {N_ITEMS} câu đánh giá theo thang điểm và 01 câu ước tính tỷ trọng — "
           f"thời gian hoàn thành khoảng 7 phút.", 11, align=J, after=3)
    P(doc, "Toàn bộ thông tin Quý Doanh nghiệp cung cấp được giữ bí mật tuyệt đối, không nêu danh tính doanh "
           "nghiệp hay người trả lời ở bất kỳ khâu nào, và chỉ sử dụng cho mục đích phân tích tổng hợp.",
      11, align=J, after=10)

    # ---------- SANG LOC
    P(doc, "PHẦN A: CÂU HỎI SÀNG LỌC", 12, True, color=NAVY, after=4)
    P(doc, f"{SCREEN[0]}. {SCREEN[1]}", 11, True, after=2)
    for o in SCREEN[2]:
        P(doc, "      [   ]   " + o, 11, after=1)
    doc.paragraphs[-1].paragraph_format.space_after = Pt(10)

    # ---------- PHAN I
    P(doc, "PHẦN B: THÔNG TIN CHUNG VỀ DOANH NGHIỆP", 12, True, color=NAVY, after=4)
    for code, var, q, opts in PART_I:
        P(doc, f"{code}. {q}", 11, True, after=2)
        for o in opts:
            P(doc, "      [   ]   " + o, 11, after=1)
        doc.paragraphs[-1].paragraph_format.space_after = Pt(7)

    # ---------- PHAN II
    P(doc, "PHẦN C: ĐÁNH GIÁ CHẤT LƯỢNG DỊCH VỤ BẢO LÃNH", 12, True, color=NAVY, after=3)
    P(doc, "Đề nghị Quý Doanh nghiệp đánh giá dựa trên trải nghiệm giao dịch bảo lãnh tại VietinBank "
           "trong 12 tháng gần đây.", 10.5, False, True, align=J, after=3)
    P(doc, "Quy ước:  1 – Hoàn toàn không đồng ý   |   2 – Không đồng ý   |   3 – Trung lập   |   "
           "4 – Đồng ý   |   5 – Hoàn toàn đồng ý", 10, False, True, after=6)

    nrow = 1 + sum(1 + len(c[3]) for c in CONSTRUCTS) + 2 + len(DEC[2])
    has_na = any(c[2] for c in CONSTRUCTS)
    ncol = 8 if has_na else 7
    tb = doc.add_table(rows=nrow, cols=ncol)
    tb.style = 'Table Grid'; tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    W = [0.72, 3.55, 0.28, 0.28, 0.28, 0.28, 0.28, 0.55]

    def row(i, vals, bold=False, fill=None, merge=False):
        r = tb.rows[i]
        if merge:
            c0 = r.cells[0]
            for k in range(1, ncol): c0 = c0.merge(r.cells[k])
            c0.text = ""
            run = c0.paragraphs[0].add_run(vals)
            run.bold = True; run.font.size = Pt(10.5); run.font.color.rgb = NAVY
            if fill: shade(c0, fill)
            return
        for k, v in enumerate(vals):
            cell = r.cells[k]; cell.width = Inches(W[k]); cell.text = ""
            p = cell.paragraphs[0]
            if k >= 2: p.alignment = C
            run = p.add_run(v); run.bold = bold
            run.font.size = Pt(8.5 if k >= 2 else 10)
            if fill: shade(cell, fill)

    head = ["Mã", "Nội dung phát biểu", "1", "2", "3", "4", "5"] + (["Chưa\ndùng"] if has_na else [])
    i = 0; row(i, head, bold=True, fill=GREY); i += 1
    for var, title, na, items in CONSTRUCTS:
        row(i, f"{title}   ({var})", fill=LIGHT, merge=True); i += 1
        for code, text in items:
            cells = [code, text] + ["(  )"] * 5 + (["(  )" if na else "—"] if has_na else [])
            row(i, cells); i += 1
    row(i, f"{DEC[1]}   ({DEC[0]})", fill=GREY, merge=True); i += 1
    row(i, "So sánh với các ngân hàng khác mà Doanh nghiệp đã hoặc đang cân nhắc sử dụng dịch vụ bảo lãnh:",
        fill=LIGHT, merge=True); i += 1
    for code, text in DEC[2]:
        row(i, [code, text] + ["(  )"] * 5 + (["—"] if has_na else [])); i += 1

    P(doc, "Ghi chú: cột “Chưa dùng” chỉ áp dụng cho nhóm III (bảo lãnh điện tử eFAST). Nếu Doanh nghiệp chưa "
           "từng sử dụng kênh này, đề nghị đánh dấu vào cột đó thay vì chọn điểm 1–5.",
      9.5, False, True, align=J, after=10)

    # ---------- PHAN D
    P(doc, "PHẦN D: TỶ TRỌNG GIAO DỊCH", 12, True, color=NAVY, after=4)
    P(doc, f"{VALIDATION[0]}. {VALIDATION[2]}", 11, True, after=2)
    for o in VALIDATION[3]:
        P(doc, "      [   ]   " + o, 11, after=1)

    doc.add_paragraph()
    P(doc, "XIN CHÂN THÀNH CẢM ƠN QUÝ DOANH NGHIỆP ĐÃ HOÀN THÀNH PHIẾU KHẢO SÁT!",
      11.5, True, align=C, color=NAVY)

    doc.save("Phieu_Khao_Sat_Chinh_Thuc_VietinBank.docx")
    print(f"Da tao phieu | {len(CONSTRUCTS)} bien x 4 cau + {len(DEC[2])} cau DEC = {N_ITEMS} cau Likert")


if __name__ == "__main__":
    build()
