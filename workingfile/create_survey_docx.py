# -*- coding: utf-8 -*-
"""
Sinh file Word Phieu Khao Sat Chinh Thuc - VietinBank
Mo hinh 7 bien doc lap + DEC (25 cau Likert 1-5)
Ma bien trung khop 100% voi ma nhan to trong mo hinh hoi quy.
"""
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NAVY = RGBColor(0x00, 0x33, 0x66)
GREY = "D9D9D9"
LIGHT = "F2F2F2"


def shade(cell, hexfill):
    tcPr = cell._element.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hexfill)
    tcPr.append(shd)


def para(doc, text, size=11.5, bold=False, italic=False, align=None, color=None, space_after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    return p


# ---------------------------------------------------------------- noi dung
GROUPS = [
    ("I. Tính cạnh tranh của biểu phí bảo lãnh (COST_COMP)", [
        ("COMP1", "Mức phí phát hành bảo lãnh tại VietinBank là hợp lý so với chất lượng dịch vụ nhận được."),
        ("COMP2", "Biểu phí bảo lãnh của VietinBank có tính cạnh tranh hơn so với các ngân hàng khác mà Doanh nghiệp đã tìm hiểu."),
        ("COMP3", "VietinBank có chính sách ưu đãi và chiết khấu phí linh hoạt cho khách hàng giao dịch thường xuyên."),
    ]),
    ("II. Tốc độ xử lý và thủ tục phát hành (PROC_SPEED)", [
        ("SPEED1", "Thời gian thẩm định và phê duyệt hạn mức bảo lãnh tại VietinBank nhanh chóng."),
        ("SPEED2", "Thủ tục hồ sơ đề nghị cấp bảo lãnh tại VietinBank đơn giản, không rườm rà."),
        ("SPEED3", "Thời gian từ khi nộp đủ hồ sơ đến khi nhận Thư bảo lãnh đáp ứng kịp tiến độ hợp đồng."),
    ]),
    ("III. Tiện ích chuyển đổi số và bảo lãnh điện tử eFAST (DIGITAL_CONV)", [
        ("DIGI1", "Doanh nghiệp nộp được đề nghị cấp bảo lãnh trực tuyến 24/7 qua nền tảng VietinBank eFAST."),
        ("DIGI2", "VietinBank phát hành Thư bảo lãnh điện tử có chữ ký số nhanh chóng, không cần bản giấy."),
        ("DIGI3", "Việc tra cứu, xác thực và theo dõi trạng thái thư bảo lãnh trực tuyến thuận tiện và đáng tin cậy."),
    ]),
    ("IV. Uy tín và thương hiệu ngân hàng (BANK_REP)", [
        ("REPU1", "VietinBank có uy tín và thương hiệu thuộc nhóm dẫn đầu thị trường ngân hàng Việt Nam."),
        ("REPU2", "Thư bảo lãnh do VietinBank phát hành được Chủ đầu tư và Bên mời thầu chấp nhận rộng rãi hơn so với thư của nhiều ngân hàng khác."),
        ("REPU3", "Năng lực tài chính của VietinBank tạo thêm uy tín cho Doanh nghiệp khi tham gia đấu thầu và ký kết hợp đồng."),
    ]),
    ("V. Quan hệ tín dụng và chính sách hạn mức (RELATIONSHIP)", [
        ("RELA1", "Lịch sử quan hệ tín dụng lâu năm với VietinBank giúp Doanh nghiệp thuận lợi hơn khi đề nghị cấp bảo lãnh."),
        ("RELA2", "Việc sử dụng đồng thời nhiều dịch vụ khác tại VietinBank (tài khoản, thanh toán, trả lương, tài trợ vốn) mang lại ưu đãi cho hoạt động bảo lãnh."),
        ("RELA3", "VietinBank cấp và điều chỉnh hạn mức bảo lãnh linh hoạt theo nhu cầu thực tế của Doanh nghiệp."),
    ]),
    ("VI. Năng lực chuyên môn và tư vấn pháp lý của cán bộ (STAFF_QUAL)", [
        ("STAFF1", "Cán bộ quan hệ khách hàng của VietinBank có chuyên môn vững về nghiệp vụ bảo lãnh."),
        ("STAFF2", "Cán bộ VietinBank tư vấn được điều khoản thư bảo lãnh phù hợp quy định pháp luật hiện hành (Luật Đấu thầu, Thông tư 61/2024/TT-NHNN) và thông lệ quốc tế URDG 758."),
        ("STAFF3", "Cán bộ VietinBank hỗ trợ kịp thời, tận tình khi Doanh nghiệp phát sinh vướng mắc liên quan đến bảo lãnh."),
    ]),
    ("VII. Chính sách ký quỹ và tài sản bảo đảm (COLL_POLICY)", [
        ("COLL1", "VietinBank áp dụng tỷ lệ ký quỹ linh hoạt, có chính sách giảm hoặc miễn ký quỹ cho khách hàng đủ điều kiện."),
        ("COLL2", "VietinBank chấp nhận đa dạng loại tài sản bảo đảm (bất động sản, máy móc thiết bị, quyền đòi nợ từ hợp đồng)."),
        ("COLL3", "Thủ tục định giá và nhận tài sản bảo đảm tại VietinBank nhanh gọn, thuận tiện."),
    ]),
]

DEC_ITEMS = [
    ("DEC1", "Khi phát sinh nhu cầu bảo lãnh, Doanh nghiệp ưu tiên lựa chọn VietinBank trước các ngân hàng khác."),
    ("DEC2", "Doanh nghiệp dành phần lớn giá trị và số lượng thư bảo lãnh của mình cho VietinBank thay vì các ngân hàng khác."),
    ("DEC3", "Trong các gói thầu và hợp đồng sắp tới, Doanh nghiệp dự định tiếp tục chọn VietinBank thay vì chuyển sang ngân hàng khác."),
    ("DEC4", "Doanh nghiệp sẵn sàng giới thiệu VietinBank cho đối tác, nhà thầu liên danh như ngân hàng phát hành bảo lãnh tốt nhất."),
]

PART_I = [
    ("1. Loại hình sở hữu của Doanh nghiệp:",
     ["Doanh nghiệp tư nhân / Công ty TNHH", "Công ty cổ phần (ngoài nhà nước)",
      "Doanh nghiệp nhà nước / có vốn nhà nước chi phối", "Doanh nghiệp có vốn đầu tư nước ngoài (FDI)",
      "Loại hình khác"]),
    ("2. Quy mô doanh thu năm gần nhất:",
     ["Dưới 20 tỷ VNĐ", "Từ 20 đến dưới 100 tỷ VNĐ", "Từ 100 đến dưới 500 tỷ VNĐ", "Từ 500 tỷ VNĐ trở lên"]),
    ("3. Thâm niên hoạt động của Doanh nghiệp:",
     ["Dưới 3 năm", "Từ 3 đến dưới 5 năm", "Từ 5 đến dưới 10 năm", "Từ 10 năm trở lên"]),
    ("4. Loại bảo lãnh Doanh nghiệp sử dụng NHIỀU NHẤT tại VietinBank (chỉ chọn 01 phương án):",
     ["Bảo lãnh dự thầu (TG)", "Bảo lãnh thực hiện hợp đồng (PG)", "Bảo lãnh tạm ứng (APG)",
      "Bảo lãnh thanh toán (BG)", "Loại bảo lãnh khác (bảo hành, tái bảo lãnh...)"]),
    ("5. Hiện Doanh nghiệp đang sử dụng dịch vụ bảo lãnh tại bao nhiêu ngân hàng?",
     ["Chỉ duy nhất VietinBank", "02 ngân hàng", "03 ngân hàng", "Từ 04 ngân hàng trở lên"]),
    ("6. Chức danh của Người đại diện trả lời phiếu:",
     ["Ban Giám đốc / CFO", "Kế toán trưởng / Trưởng phòng Tài chính",
      "Trưởng phòng Đấu thầu / Mua hàng", "Chuyên viên phụ trách bảo lãnh"]),
]


def build():
    doc = docx.Document()
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Inches(1.0)
        s.left_margin = s.right_margin = Inches(1.0)

    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(11.5)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')

    C = WD_ALIGN_PARAGRAPH.CENTER
    J = WD_ALIGN_PARAGRAPH.JUSTIFY

    para(doc, "NGÂN HÀNG TMCP CÔNG THƯƠNG VIỆT NAM (VIETINBANK)", 12, True, align=C, color=NAVY, space_after=0)
    para(doc, "KHẢO SÁT Ý KIẾN KHÁCH HÀNG DOANH NGHIỆP VỀ DỊCH VỤ BẢO LÃNH", 11, True, True, align=C, color=NAVY)
    para(doc, "PHIẾU KHẢO SÁT CHÍNH THỨC", 15, True, align=C, color=NAVY, space_after=2)
    para(doc, "Các yếu tố ảnh hưởng đến quyết định lựa chọn dịch vụ bảo lãnh ngân hàng "
              "của khách hàng doanh nghiệp tại VietinBank", 11.5, False, True, align=C, space_after=12)

    para(doc, "Kính gửi Quý Doanh nghiệp,", 11.5, True, space_after=4)
    para(doc, "Nhằm nâng cao chất lượng dịch vụ bảo lãnh và thiết kế các gói sản phẩm phù hợp hơn với nhu cầu "
              "của Quý Doanh nghiệp, chúng tôi tiến hành cuộc khảo sát này. Kính mong Quý Doanh nghiệp dành ít "
              "phút đưa ra đánh giá khách quan. Phiếu gồm 06 câu thông tin chung và 25 câu đánh giá theo thang điểm, tổng cộng 31 câu, thời gian hoàn thành khoảng 5–7 phút.",
         11.5, align=J, space_after=4)
    para(doc, "Toàn bộ thông tin Quý Doanh nghiệp cung cấp được giữ bí mật tuyệt đối và chỉ sử dụng cho mục đích "
              "phân tích tổng hợp phục vụ nghiên cứu. Xin chân thành cảm ơn sự hợp tác của Quý Doanh nghiệp.",
         11.5, align=J, space_after=12)

    para(doc, "PHẦN I: THÔNG TIN CHUNG VỀ DOANH NGHIỆP", 12.5, True, color=NAVY, space_after=8)
    for q, opts in PART_I:
        para(doc, q, 11.5, True, space_after=2)
        for o in opts:
            para(doc, "      [   ]   " + o, 11.5, space_after=1)
        doc.paragraphs[-1].paragraph_format.space_after = Pt(8)

    para(doc, "PHẦN II: ĐÁNH GIÁ CÁC YẾU TỐ (THANG ĐO LIKERT 5 MỨC ĐỘ)", 12.5, True, color=NAVY, space_after=4)
    para(doc, "Quy ước:  1 – Hoàn toàn không đồng ý   |   2 – Không đồng ý   |   3 – Trung lập   |   "
              "4 – Đồng ý   |   5 – Hoàn toàn đồng ý", 11, False, True, space_after=8)

    rows = 1
    for _, items in GROUPS:
        rows += 1 + len(items)
    rows += 2 + len(DEC_ITEMS)

    tb = doc.add_table(rows=rows, cols=7)
    tb.style = 'Table Grid'
    tb.alignment = WD_TABLE_ALIGNMENT.CENTER
    widths = [Inches(0.75), Inches(3.85), Inches(0.3), Inches(0.3), Inches(0.3), Inches(0.3), Inches(0.3)]

    def setrow(i, vals, bold=False, fill=None, merge2=False):
        r = tb.rows[i]
        if merge2:
            c = r.cells[0]
            for k in range(1, 7):
                c = c.merge(r.cells[k])
            c.text = ""
            p = c.paragraphs[0]
            run = p.add_run(vals)
            run.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = NAVY
            if fill:
                shade(c, fill)
            return
        for k, v in enumerate(vals):
            cell = r.cells[k]
            cell.width = widths[k]
            cell.text = ""
            p = cell.paragraphs[0]
            if k >= 2:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(v)
            run.bold = bold
            run.font.size = Pt(10.5 if k >= 2 else 11)
            if fill:
                shade(cell, fill)

    i = 0
    setrow(i, ["Mã biến", "Nội dung phát biểu đánh giá", "1", "2", "3", "4", "5"], bold=True, fill=GREY)
    i += 1

    for title, items in GROUPS:
        setrow(i, title, fill=LIGHT, merge2=True); i += 1
        for code, text in items:
            setrow(i, [code, text, "(  )", "(  )", "(  )", "(  )", "(  )"]); i += 1

    setrow(i, "VIII. MỨC ĐỘ ƯU TIÊN LỰA CHỌN VIETINBANK (BIẾN PHỤ THUỘC — DEC)", fill=GREY, merge2=True); i += 1
    setrow(i, "So sánh với các ngân hàng khác mà Doanh nghiệp đã hoặc đang cân nhắc sử dụng dịch vụ bảo lãnh:",
           fill=LIGHT, merge2=True); i += 1
    for code, text in DEC_ITEMS:
        setrow(i, [code, text, "(  )", "(  )", "(  )", "(  )", "(  )"]); i += 1

    doc.add_paragraph()
    para(doc, "XIN CHÂN THÀNH CẢM ƠN QUÝ DOANH NGHIỆP ĐÃ HOÀN THÀNH PHIẾU KHẢO SÁT!",
         12, True, align=WD_ALIGN_PARAGRAPH.CENTER, color=NAVY)

    doc.save("Phieu_Khao_Sat_Chinh_Thuc_VietinBank.docx")
    n = sum(len(x[1]) for x in GROUPS) + len(DEC_ITEMS)
    print("Da tao Phieu_Khao_Sat_Chinh_Thuc_VietinBank.docx |", len(GROUPS), "nhom x 3 +", len(DEC_ITEMS), "DEC =", n, "cau")


if __name__ == "__main__":
    build()
